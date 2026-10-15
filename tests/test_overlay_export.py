"""Overlays outside the app: `render --overlay` and the exports (#215)."""

from __future__ import annotations

import json
import os
from pathlib import Path

from PySide6.QtGui import QColor, QImage
from PySide6.QtWidgets import QApplication

from grafli import theme
from grafli.format import parse
from grafli.view import GrafliView

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_BOARD = """\
#!grafli v2
@ box a "A" 0,0 120x60
@ box b "B" 300,0 120x60
@ box c "C" 600,0 120x60
@ arrow a -> b ""
@ arrow b -> c ""
"""

TESTS = {
    "title": "Test results",
    "producer": "overlay_tests.py",
    "kind": "category",
    "categories": [
        {"id": "pass", "label": "Passing", "color": "#00ff00"},
        {"id": "fail", "label": "Failing", "color": "#ff0000"},
    ],
    "entries": {"a": {"category": "fail"}, "b": {"category": "pass"}},
}


def _map(tmp_path: Path, src: str = _BOARD, **overlays) -> Path:
    board = tmp_path / "map.grafli"
    board.write_text(src)
    res = tmp_path / "map-res"
    res.mkdir(exist_ok=True)
    for name, data in overlays.items():
        (res / f"{name}.overlay.json").write_text(
            data if isinstance(data, str) else json.dumps(data))
    return board


def _render(board: Path, out: Path, *flags: str) -> int:
    from grafli.app import _cmd_render
    return _cmd_render([str(board), str(out), *flags])


def _view_of(board: Path) -> GrafliView:
    QApplication.instance() or QApplication([])
    view = GrafliView()
    view._file_path = board      # the view is its own window here
    view.load_board(parse(board.read_text()))
    return view


# ── render --overlay ────────────────────────────────────────────────

def test_render_overlay_colours_the_board_and_adds_its_legend(tmp_path):
    board = _map(tmp_path, tests=TESTS)
    assert _render(board, tmp_path / "plain.png") == 0
    assert _render(board, tmp_path / "tests.png", "--overlay", "tests") == 0
    plain = QImage(str(tmp_path / "plain.png"))
    shown = QImage(str(tmp_path / "tests.png"))
    assert shown.width() >= plain.width()
    assert shown.height() > plain.height()
    # The page takes the overlay's darkened canvas, the legend band too.
    heat_bg = theme.HEATMAP_BG.name()
    assert QColor(shown.pixel(2, 2)).name() == heat_bg
    assert QColor(shown.pixel(2, shown.height() - 3)).name() == heat_bg
    assert QColor(plain.pixel(2, 2)).name() == theme.SCENE_BG.name()


def test_render_overlay_to_svg_names_the_legend(tmp_path):
    board = _map(tmp_path, tests=TESTS)
    assert _render(board, tmp_path / "plain.svg") == 0
    assert _render(board, tmp_path / "tests.svg", "--overlay", "tests") == 0
    plain = (tmp_path / "plain.svg").read_text()
    shown = (tmp_path / "tests.svg").read_text()
    for text in ("Test results", "overlay_tests.py", "Passing", "Failing",
                 "no data"):
        assert text not in plain
        assert text in shown


def test_render_overlay_with_the_palette_legend_draws_both(tmp_path):
    src = (_BOARD.replace("v2", "v4")
           .replace('"A" 0,0 120x60', '"A" 0,0 120x60 %core'))
    board = _map(tmp_path, src, tests=TESTS)
    (tmp_path / "map-res" / "palette.json").write_text(json.dumps(
        {"entries": {"core": {"color": "#3366cc", "label": "Core model"}}}))
    assert _render(board, tmp_path / "both.svg", "--overlay", "tests",
                   "--legend") == 0
    svg = (tmp_path / "both.svg").read_text()
    assert "Core model" in svg
    assert "Test results" in svg


def test_render_overlay_names_the_available_ones_when_unknown(tmp_path,
                                                               capsys):
    board = _map(tmp_path, tests=TESTS)
    assert _render(board, tmp_path / "x.png", "--overlay", "churn") == 2
    err = capsys.readouterr().err
    assert "no overlay 'churn'" in err
    assert "available: tests" in err
    assert not (tmp_path / "x.png").exists()


def test_render_overlay_reports_a_broken_file(tmp_path, capsys):
    board = _map(tmp_path, tests="{ not json")
    assert _render(board, tmp_path / "x.png", "--overlay", "tests") == 2
    assert "--overlay" in capsys.readouterr().err


def test_render_without_overlay_is_unchanged_by_overlay_files(tmp_path):
    board = _map(tmp_path)
    assert _render(board, tmp_path / "before.png") == 0
    _map(tmp_path, tests=TESTS)
    assert _render(board, tmp_path / "after.png") == 0
    assert (QImage(str(tmp_path / "before.png"))
            == QImage(str(tmp_path / "after.png")))


def test_the_app_svg_export_carries_the_shown_overlay(tmp_path):
    view = _view_of(_map(tmp_path, tests=TESTS))
    assert b"Test results" not in bytes(view._render_svg_bytes())
    view.show_overlay("tests")
    assert b"Test results" in bytes(view._render_svg_bytes())
    view._clear_complexity_heatmap()
    assert b"Test results" not in bytes(view._render_svg_bytes())


# ── the PDF/PPTX exports ────────────────────────────────────────────

_FLOW_BOARD = _BOARD + """\
@ note n 0,200 "a remark"
@ bookmark all "All of it" @a,b,c,n "Where the tests stand."
@ flow review "Review" all
"""


def _pdf_text(path: Path) -> str:
    from PySide6.QtPdf import QPdfDocument
    doc = QPdfDocument()
    doc.load(str(path))
    text = "\n".join(doc.getAllText(i).text() for i in range(doc.pageCount()))
    doc.close()
    return text


def _export(board: Path, out: Path, *flags: str) -> int:
    from grafli.app import _cmd_export
    return _cmd_export([str(board), str(out), *flags])


def test_a_pdf_exported_from_the_app_with_an_overlay_on_shows_it(
        tmp_path, monkeypatch):
    QApplication.instance() or QApplication([])
    from PySide6.QtWidgets import QFileDialog
    from grafli.app import MainWindow
    board = _map(tmp_path, _FLOW_BOARD, tests=TESTS)
    w = MainWindow(str(board))
    monkeypatch.setattr(w._view, "toast", lambda *a, **k: None)
    outs = iter([tmp_path / "plain.pdf", tmp_path / "shown.pdf"])
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(next(outs)), "")))
    w._export_flow_pdf()
    w._view.show_overlay("tests")
    w._export_flow_pdf()
    plain = _pdf_text(tmp_path / "plain.pdf")
    shown = _pdf_text(tmp_path / "shown.pdf")
    for text in ("Test results", "Passing", "Failing", "no data"):
        assert text not in plain
        assert text in shown
    # The export leaves the overlay on the canvas as it found it.
    assert w._view._overlay_shown() is not None
    assert (w._view._scene.backgroundBrush().color().name()
            == theme.HEATMAP_BG.name())


def test_export_overlay_carries_it_into_the_pdf(tmp_path):
    board = _map(tmp_path, _FLOW_BOARD, tests=TESTS)
    assert _export(board, tmp_path / "plain.pdf") == 0
    assert _export(board, tmp_path / "shown.pdf", "--overlay", "tests") == 0
    assert "Test results" not in _pdf_text(tmp_path / "plain.pdf")
    assert "Test results" in _pdf_text(tmp_path / "shown.pdf")


def test_export_overlay_adds_the_legend_card_to_each_pptx_diagram_slide(
        tmp_path):
    from pptx import Presentation
    board = _map(tmp_path, _FLOW_BOARD, tests=TESTS)
    assert _export(board, tmp_path / "plain.pptx") == 0
    assert _export(board, tmp_path / "shown.pptx", "--overlay", "tests") == 0

    def pictures(path):
        return [sum(1 for s in slide.shapes if s.shape_type == 13)
                for slide in Presentation(str(path)).slides]

    plain, shown = pictures(tmp_path / "plain.pptx"), pictures(
        tmp_path / "shown.pptx")
    assert shown[0] == plain[0]                  # the cover has no diagram
    assert shown[1] == plain[1] + 1              # the card beside it


def test_export_overlay_names_the_available_ones_when_unknown(tmp_path,
                                                               capsys):
    board = _map(tmp_path, _FLOW_BOARD, tests=TESTS)
    assert _export(board, tmp_path / "x.pdf", "--overlay", "churn") == 2
    assert "available: tests" in capsys.readouterr().err
    assert not (tmp_path / "x.pdf").exists()


def test_a_shown_overlay_keeps_the_notes_in_the_slide_raster(tmp_path):
    from grafli.slideplan import build_slide_plan, live_overlays
    view = _view_of(_map(tmp_path, _FLOW_BOARD, tests=TESTS))
    plan = build_slide_plan(view, view._board.flow_by_id("review"))[1]
    assert [it.note.id for it in live_overlays(view, plan)] == ["n"]
    view.show_overlay("tests")
    assert live_overlays(view, plan) == []
