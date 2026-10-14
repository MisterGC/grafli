"""`!ink` (#204): a box's label and icon take their ink from its fill and drop
the light plate; a plain note's text takes its %color. Without `!ink`, a
box keeps the plate and a coloured note keeps its note-blue text."""

from __future__ import annotations

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QRectF, Qt
from PySide6.QtGui import QColor, QImage, QPainter
from PySide6.QtWidgets import QApplication

from grafli import theme
from grafli.format import parse


def _app() -> QApplication:
    return QApplication.instance() or QApplication([])


def _view(text: str):
    from grafli.view import GrafliView
    _app()
    view = GrafliView()
    view.load_board(parse(text))
    return view


def _render(view, rect: QRectF) -> QImage:
    """The scene inside *rect* at 1:1, on the canvas colour."""
    view.scene().clearSelection()
    img = QImage(int(rect.width()), int(rect.height()),
                 QImage.Format.Format_ARGB32)
    img.fill(QColor(theme.SCENE_BG))
    painter = QPainter(img)
    view.scene().render(painter, QRectF(img.rect()), rect)
    painter.end()
    return img


def _luminance(c: QColor) -> float:
    def lin(v: float) -> float:
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    return (0.2126 * lin(c.redF()) + 0.7152 * lin(c.greenF())
            + 0.0722 * lin(c.blueF()))


def _contrast(a: QColor, b: QColor) -> float:
    la, lb = sorted((_luminance(a), _luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


BOXES = ('#!grafli v4\n'
         '@ box ink "Dark" 0,0 200x80 %primary !ink\n'
         '@ box plate "Dark" 0,120 200x80 %primary\n')


def test_a_dark_ink_box_draws_a_light_label_that_reads_on_its_fill():
    view = _view(BOXES)
    ink, plate = view._box_items["ink"], view._box_items["plate"]
    fill = ink.brush().color()
    label = ink._label.defaultTextColor()
    assert _luminance(label) > _luminance(fill)
    assert _contrast(label, fill) >= 4.5
    # Without !ink the label keeps the board ink on its light plate.
    assert plate._label.defaultTextColor() == QColor(theme.INK)


def _plate_corner(item) -> QRectF:
    """A sliver inside the label plate's top-left corner, clear of glyphs."""
    r = item._label.sceneBoundingRect()
    return QRectF(r.left() + 1, r.top() + 1, 3, 3)


def test_a_dark_ink_box_draws_no_plate_behind_its_label():
    view = _view(BOXES)
    fill = view._box_items["ink"].brush().color()
    for bid, has_plate in (("ink", False), ("plate", True)):
        item = view._box_items[bid]
        px = _render(view, _plate_corner(item)).pixelColor(1, 1)
        if has_plate:
            assert _luminance(px) > _luminance(fill) + 0.1
        else:
            assert px.rgb() == fill.rgb()


def test_a_dark_ink_box_draws_its_icon_in_the_label_ink(monkeypatch):
    from grafli import items
    seen = []
    real = items.iconset.paint_icon

    def spy(painter, name, rect, color):
        seen.append(QColor(color))
        real(painter, name, rect, color)

    monkeypatch.setattr(items.iconset, "paint_icon", spy)
    view = _view('#!grafli v4\n'
                 '@ box b "Dark" 0,0 200x80 %primary !ink *lead:star\n')
    item = view._box_items["b"]
    _render(view, item.sceneBoundingRect())
    assert seen and all(c == item._label.defaultTextColor() for c in seen)
    assert _contrast(seen[0], item.brush().color()) >= 4.5


def test_toggling_ink_on_a_box_switches_the_label_ink():
    view = _view(BOXES)
    item = view._box_items["plate"]
    item.set_ink(True)
    assert item._label.defaultTextColor() != QColor(theme.INK)
    item.set_ink(False)
    assert item._label.defaultTextColor() == QColor(theme.INK)


NOTES = ('#!grafli v4\n'
         '@ note rose 0,0 "rose text" %rose !ink\n'
         '@ note muted 0,100 "muted text" %muted !ink\n'
         '@ note blue 0,200 "blue text" %rose\n')


def _count_near(img: QImage, target: QColor, tol: int = 24) -> int:
    n = 0
    for y in range(img.height()):
        for x in range(img.width()):
            c = img.pixelColor(x, y)
            if (abs(c.red() - target.red()) <= tol
                    and abs(c.green() - target.green()) <= tol
                    and abs(c.blue() - target.blue()) <= tol):
                n += 1
    return n


def test_rose_and_muted_ink_notes_draw_their_text_in_those_inks():
    view = _view(NOTES)
    blue = QColor(theme.NOTE_PEN_COLOR)
    for nid, token in (("rose", "%rose"), ("muted", "%muted")):
        item = view._note_items[nid]
        want = QColor(theme.resolve_color(token))
        assert item._parse_note()[2] == want
        img = _render(view, item.sceneBoundingRect())
        assert _count_near(img, want, tol=6) > 20
        assert _count_near(img, blue, tol=6) == 0


def test_a_coloured_note_without_ink_keeps_note_blue_text():
    view = _view(NOTES)
    item = view._note_items["blue"]
    assert item._parse_note()[2] == theme.NOTE_PEN_COLOR


def test_a_coloured_note_without_ink_renders_0_px_different():
    # %color without !ink only tints a note's icon: the text note renders
    # exactly as the same note with no colour at all.
    coloured = _view('#!grafli v1\n@ note n 0,0 "blue text" %rose\n')
    plain = _view('#!grafli v1\n@ note n 0,0 "blue text"\n')
    rect = coloured._note_items["n"].sceneBoundingRect()
    assert rect == plain._note_items["n"].sceneBoundingRect()
    assert _render(coloured, rect) == _render(plain, rect)


def test_toggling_ink_on_a_note_switches_its_text_colour():
    view = _view(NOTES)
    item = view._note_items["blue"]
    item.set_ink(True)
    assert item._parse_note()[2] == QColor(theme.resolve_color("%rose"))
    item.set_ink(False)
    assert item._parse_note()[2] == theme.NOTE_PEN_COLOR


# ── !ink note text in the PDF and PPTX exports (#206) ───────────

INK_TOUR = ('#!grafli v4\n'
            '@ box a "A" 0,0 160x60\n'
            '@ note n 0,100 "Rose ink text here" %rose !ink\n'
            '@ bookmark bm "Overlay" @a,n\n'
            '@ bookmark bt "Text" @n\n'
            '@ flow tour "Tour" bm bt\n')


def _pdf_pages(tmp_path, text: str, name: str) -> list[QImage]:
    """The PDF export's content slides (the cover left out) as images."""
    from PySide6.QtCore import QSize
    from PySide6.QtPdf import QPdfDocument
    from grafli.pdfexport import export_flow_to_pdf
    board = parse(text)
    out = tmp_path / f"{name}.pdf"
    export_flow_to_pdf(_view(text), board.flow_by_id("tour"), out)
    doc = QPdfDocument()
    doc.load(str(out))
    pages = [doc.render(i, QSize(1600, 900)) for i in range(doc.pageCount())]
    doc.close()
    return pages[1:]


def test_an_ink_note_prints_its_text_in_its_colour_in_the_pdf(tmp_path):
    rose = QColor(theme.resolve_color("%rose"))
    overlay, text = _pdf_pages(tmp_path, INK_TOUR, "ink")
    assert _count_near(overlay, rose, tol=12) > 20      # note over the diagram
    assert _count_near(text, rose, tol=12) > 20         # text slide
    plain = _pdf_pages(tmp_path, INK_TOUR.replace(" !ink", ""), "plain")
    assert all(_count_near(p, rose, tol=12) == 0 for p in plain)


def test_a_coloured_note_without_ink_prints_0_px_different_in_the_pdf(
        tmp_path):
    coloured = _pdf_pages(tmp_path, INK_TOUR.replace(" !ink", ""), "rose")
    plain = _pdf_pages(tmp_path, INK_TOUR.replace(" %rose !ink", ""), "none")
    assert [p.size() for p in coloured] == [p.size() for p in plain]
    assert coloured == plain


def _pptx_note_runs(tmp_path, text: str, name: str) -> list:
    from pptx import Presentation
    from grafli.pptxexport import export_flow_to_pptx
    board = parse(text)
    out = tmp_path / f"{name}.pptx"
    export_flow_to_pptx(_view(text), board.flow_by_id("tour"), out)
    return [r for slide in list(Presentation(out).slides)[1:]
            for sh in slide.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs
            if "Rose ink" in r.text]


def test_an_ink_note_prints_its_text_in_its_colour_in_the_pptx(tmp_path):
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_COLOR_TYPE
    rose = QColor(theme.resolve_color("%rose"))
    runs = _pptx_note_runs(tmp_path, INK_TOUR, "ink")
    assert len(runs) == 2                               # overlay + text slide
    for run in runs:
        assert run.font.color.type == MSO_COLOR_TYPE.RGB
        assert run.font.color.rgb == RGBColor(rose.red(), rose.green(),
                                              rose.blue())


def test_a_note_without_ink_keeps_the_theme_text_colour_in_the_pptx(
        tmp_path):
    from pptx.enum.dml import MSO_COLOR_TYPE, MSO_THEME_COLOR
    runs = _pptx_note_runs(tmp_path, INK_TOUR.replace(" !ink", ""), "plain")
    assert len(runs) == 2
    for run in runs:
        assert run.font.color.type == MSO_COLOR_TYPE.SCHEME
        assert run.font.color.theme_color == MSO_THEME_COLOR.TEXT_1
        assert run.font.name == "+mn-lt"


def test_only_a_plain_ink_note_takes_its_colour_in_the_exports():
    from grafli.items import note_ink
    for text in ('md:\n**bold**', 'T: a task', 'code: x = 1',
                 'Ann: one\nBob: two'):
        note = parse(f'#!grafli v4\n@ note n 0,0 """\n{text}\n""" %rose !ink\n'
                     ).notes[0]
        assert note_ink(note) is None, text
    note = parse('#!grafli v4\n@ note n 0,0 "plain" %rose !ink\n').notes[0]
    assert note_ink(note) == QColor(theme.resolve_color("%rose"))
