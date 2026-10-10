"""Path stops (D9): a bookmark's focus may name arrows by their ``~id=``. The
stop frames the arrow's two ends and, while a tour shows it or `grafli render
--bookmark` draws it, the arrow takes the tour accent, thicker, with
everything else dimmed (#167)."""

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication

from grafli import theme
from grafli.flows import bookmark_target_rect, emphasis_for, isolate_focus
from grafli.format import parse
from grafli.items import ArrowLineItem

BOARD = "\n".join([
    "#!grafli v3",
    '@ box client "Client" 0,0 160x60',
    '@ box server "Server" 600,300 160x60',
    '@ box db "DB" 600,600 160x60',
    '@ box aside "Aside" 300,0 160x60',
    '@ arrow client -> server "request" ~id=call',
    '@ arrow server -> db "query"',
    '@ arrow server -> client ~id=reply',
    '@ bookmark b_call "The call" @call',
    '@ bookmark b_mixed "The call and the DB" @call,db',
    '@ bookmark b_aside "Aside" @aside',
    '@ bookmark b_reply "The reply" @reply',
    '@ flow f "Request path" b_aside b_call',
]) + "\n"


def _app() -> QApplication:
    return QApplication.instance() or QApplication([])


def _view(board):
    app = _app()
    from grafli.view import GrafliView
    view = GrafliView()
    view.resize(800, 600)
    view.show()
    view.load_board(board)
    app.processEvents()
    return view


def _lines(view, arrow_id: str) -> list[ArrowLineItem]:
    arrow = view._board.arrow_by_id(arrow_id)
    return [g for g in view._arrow_items
            if isinstance(g, ArrowLineItem) and g.data(0) is arrow]


def test_the_framed_rect_covers_both_ends_of_the_arrow():
    view = _view(parse(BOARD))
    rect = bookmark_target_rect(view, view._board.bookmark_by_id("b_call"))
    assert not rect.isNull()
    for end in ("client", "server"):
        assert rect.contains(view._box_items[end].sceneBoundingRect())
    for line in _lines(view, "call"):
        assert rect.contains(line.sceneBoundingRect())
    # Nothing else widens the frame.
    assert not rect.contains(view._box_items["db"].sceneBoundingRect())


def test_an_arrow_and_an_element_frame_together():
    view = _view(parse(BOARD))
    rect = bookmark_target_rect(view, view._board.bookmark_by_id("b_mixed"))
    for end in ("client", "server", "db"):
        assert rect.contains(view._box_items[end].sceneBoundingRect())


def test_emphasis_names_the_arrows_and_keeps_their_ends():
    board = parse(BOARD)
    assert emphasis_for(board, board.bookmark_by_id("b_call")) == (
        {"call"}, {"client", "server"})
    assert emphasis_for(board, board.bookmark_by_id("b_mixed")) == (
        {"call"}, {"client", "server", "db"})
    assert emphasis_for(board, board.bookmark_by_id("b_aside")) is None


def test_the_tour_stop_emphasises_the_arrow_and_dims_the_rest():
    view = _view(parse(BOARD))
    plain_width = _lines(view, "call")[0].pen().widthF()
    view.play_flow("f")
    player = view._flow_player
    # A stop framing only boxes changes no arrow.
    assert view._present_emphasis is None
    player.next()
    for line in _lines(view, "call"):
        assert line.pen().color() == theme.FLOWS_ACCENT
        assert line.pen().widthF() == plain_width * 2
        assert line.opacity() == 1.0
    for end in ("client", "server"):
        assert view._box_items[end].opacity() == 1.0
    for other in ("db", "aside"):
        assert view._box_items[other].opacity() == 0.08
    for line in _lines(view, "reply"):
        assert line.opacity() == 0.08
    # Leaving the tour restores the arrow and the opacities.
    player.stop()
    assert view._present_emphasis is None
    for line in _lines(view, "call"):
        assert line.pen().color() != theme.FLOWS_ACCENT
        assert line.pen().widthF() == plain_width
        assert line.opacity() == 1.0
    assert view._box_items["db"].opacity() == 1.0


def test_a_reverse_arrow_merged_into_one_line_is_emphasised():
    # call and reply run between the same pair in opposite directions, so
    # they draw as one line carrying call; emphasising reply finds it.
    view = _view(parse(BOARD))
    board = view._board
    assert _lines(view, "reply") == []
    view._set_presentation_emphasis(
        emphasis_for(board, board.bookmark_by_id("b_reply")))
    lines = _lines(view, "call")
    assert lines
    for line in lines:
        assert line.pen().color() == theme.FLOWS_ACCENT
        assert line.opacity() == 1.0


def test_isolating_an_arrow_keeps_it_and_its_ends():
    view = _view(parse(BOARD))
    with isolate_focus(view, ["call"]):
        for end in ("client", "server"):
            assert view._box_items[end].isVisible()
        assert not view._box_items["db"].isVisible()
        assert all(line.isVisible() for line in _lines(view, "call"))
        query = [g for g in view._arrow_items
                 if g.data(0) is view._board.arrows[1]]
        assert query and not any(g.isVisible() for g in query)


def test_a_board_link_to_an_arrow_frames_it():
    view = _view(parse(BOARD))
    assert view.frame_fragment("call")
    visible = view.mapToScene(view.viewport().rect()).boundingRect()
    for end in ("client", "server"):
        assert visible.contains(view._box_items[end].sceneBoundingRect())


def test_render_bookmark_draws_the_arrow_in_the_accent(tmp_path: Path):
    from grafli.app import _cmd_render
    path = tmp_path / "path.grafli"
    path.write_text(BOARD)
    accent = theme.FLOWS_ACCENT.name().lower()
    call = tmp_path / "call.svg"
    assert _cmd_render([str(path), str(call), "--bookmark", "b_call"]) == 0
    assert accent in call.read_text().lower()
    aside = tmp_path / "aside.svg"
    assert _cmd_render([str(path), str(aside), "--bookmark", "b_aside"]) == 0
    assert accent not in aside.read_text().lower()


def test_export_check_takes_an_arrow_id_as_a_focus(tmp_path: Path, capsys):
    import json
    from grafli.app import _cmd_export
    path = tmp_path / "path.grafli"
    path.write_text(BOARD.replace("@call,db", "@call,gone"))
    assert _cmd_export([str(path), "--check", "--json"]) == 1
    report = json.loads(capsys.readouterr().out)
    assert report["dangling"] == [
        "bookmark 'b_mixed' anchors missing element 'gone'"]


# ── slide export ──


def _accent_pixels(image) -> int:
    """Pixels close to the tour accent on a rendered page."""
    accent = theme.FLOWS_ACCENT
    count = 0
    for y in range(0, image.height(), 2):
        for x in range(0, image.width(), 2):
            c = image.pixelColor(x, y)
            if (abs(c.red() - accent.red()) < 24
                    and abs(c.green() - accent.green()) < 24
                    and abs(c.blue() - accent.blue()) < 24):
                count += 1
    return count


def _pdf_pages(path: Path) -> list:
    from PySide6.QtCore import QSize
    from PySide6.QtPdf import QPdfDocument
    doc = QPdfDocument()
    doc.load(str(path))
    pages = [doc.render(i, QSize(960, 540)) for i in range(doc.pageCount())]
    doc.close()
    return pages


def test_the_slide_plan_carries_the_stop_emphasis():
    from grafli.slideplan import build_slide_plan
    view = _view(parse(BOARD))
    plans = build_slide_plan(view, view._board.flow_by_id("f"))
    # Cover, then b_aside (no arrow), then b_call.
    assert plans[1].emphasis is None
    assert plans[2].emphasis == ({"call"}, {"client", "server"})


def test_an_exported_arrow_stop_draws_the_arrow_in_the_accent(tmp_path):
    from grafli.pdfexport import export_flow_to_pdf
    view = _view(parse(BOARD))
    out = tmp_path / "path.pdf"
    export_flow_to_pdf(view, view._board.flow_by_id("f"), out)
    cover, aside, call = _pdf_pages(out)
    assert _accent_pixels(call) > 50
    assert _accent_pixels(aside) == 0
    # The export leaves the live view as it found it.
    assert view._present_emphasis is None
    for line in _lines(view, "call"):
        assert line.pen().color() != theme.FLOWS_ACCENT


def test_pptx_export_of_an_arrow_stop_restores_the_view(tmp_path):
    from pptx import Presentation
    from grafli.pptxexport import export_flow_to_pptx
    view = _view(parse(BOARD))
    out = tmp_path / "path.pptx"
    export_flow_to_pptx(view, view._board.flow_by_id("f"), out)
    assert len(Presentation(str(out)).slides) == 3
    assert view._present_emphasis is None
