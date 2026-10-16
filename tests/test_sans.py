"""`!sans` (#205, #206): a box's label (and its doc line) and a note's text
are set in the bundled proportional face, Inter, on the canvas, in `grafli
render` and in the PDF, PPTX and HTML exports. Without `!sans` a box keeps
grafli's mono face and a note its handwritten or mono one."""

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtGui import QFontInfo, QImage
from PySide6.QtWidgets import QApplication

from grafli.constants import FONT_FAMILY, SANS_FONT_FAMILY
from grafli.format import parse


def _app() -> QApplication:
    app = QApplication.instance() or QApplication([])
    from grafli.fonts import register_bundled_fonts
    register_bundled_fonts()
    return app


def _view(text: str):
    from grafli.view import GrafliView
    _app()
    view = GrafliView()
    view.load_board(parse(text))
    return view


BOARD = ('#!grafli v4\n'
         '@ box s "Sans title" 0,0 240x80 %primary !sans\n'
         '@ box m "Mono title" 0,120 240x80 %primary\n')


def test_a_sans_box_label_resolves_to_inter_and_a_plain_one_to_mono():
    view = _view(BOARD)
    sans = view._box_items["s"]._label.font()
    mono = view._box_items["m"]._label.font()
    assert QFontInfo(sans).family() == SANS_FONT_FAMILY
    assert QFontInfo(mono).family() == FONT_FAMILY


def test_sans_keeps_the_size_and_emphasis():
    view = _view(BOARD.replace("%primary !sans", "%primary ~large !sans !bold"))
    font = view._box_items["s"]._label.font()
    assert QFontInfo(font).family() == SANS_FONT_FAMILY
    assert font.bold()
    plain = view._box_items["m"]._label.font()
    assert font.pointSizeF() > plain.pointSizeF()


def test_a_sans_box_doc_line_takes_the_sans_face():
    view = _view(BOARD)
    assert view._box_items["s"]._doc_line_font().family() == SANS_FONT_FAMILY
    assert view._box_items["m"]._doc_line_font().family() == FONT_FAMILY


def test_set_sans_switches_the_label_and_back():
    view = _view(BOARD)
    item = view._box_items["m"]
    mono = item._label.font()
    item.set_sans(True)
    assert item.box.sans
    assert item._label.font().family() == SANS_FONT_FAMILY
    item.set_sans(False)
    assert item._label.font() == mono


def _render(tmp_path: Path, name: str, text: str) -> QImage:
    from grafli.app import _cmd_render
    board = tmp_path / f"{name}.grafli"
    board.write_text(text)
    out = tmp_path / f"{name}.png"
    assert _cmd_render([str(board), str(out)]) == 0
    return QImage(str(out))


def test_render_draws_a_sans_box_differently_from_mono(tmp_path: Path):
    _app()
    one = '#!grafli v4\n@ box s "Sans title" 0,0 240x80 %primary !sans\n'
    sans = _render(tmp_path, "sans", one)
    mono = _render(tmp_path, "mono", one.replace(" !sans", "")
                   .replace("v4", "v1"))
    assert sans.size() == mono.size()
    assert sans != mono


def _stop_raster(view, board, out: Path) -> QImage:
    """The board raster the PDF export paints for the flow's one stop."""
    from grafli.pdfexport import export_flow_to_pdf
    shots: list[QImage] = []
    scene = view._scene
    original = scene.render

    def capture(painter, target, source, *args):
        original(painter, target, source, *args)
        shots.append(QImage(painter.device()))

    scene.render = capture
    try:
        export_flow_to_pdf(view, board.flow_by_id("tour"), out)
    finally:
        del scene.render
    assert out.read_bytes().startswith(b"%PDF")
    assert len(shots) == 1
    return shots[0]


TOUR = ('@ bookmark bm "Title" @s\n'
        '@ flow tour "Tour" bm\n')


def test_the_pdf_export_draws_a_sans_box_in_inter(tmp_path: Path):
    text = ('#!grafli v4\n@ box s "Sans title" 0,0 240x80 %primary !sans\n'
            + TOUR)
    board = parse(text)
    view = _view(text)
    fonts: list[str] = []
    label = view._box_items["s"]._label
    original = label.paint

    def paint(painter, *args):
        fonts.append(QFontInfo(label.font()).family())
        original(painter, *args)

    label.paint = paint
    sans = _stop_raster(view, board, tmp_path / "sans.pdf")
    assert fonts and set(fonts) == {SANS_FONT_FAMILY}

    plain_text = text.replace(" !sans", "").replace("v4", "v1")
    plain = parse(plain_text)
    mono = _stop_raster(_view(plain_text), plain, tmp_path / "mono.pdf")
    assert sans.size() == mono.size()
    assert sans != mono


NOTES = ('#!grafli v4\n'
         '@ note s 0,0 "A sans note" !sans\n'
         '@ note h 0,80 "A handwritten note"\n'
         '@ note m 0,160 "A mono note" !mono\n')


def test_a_sans_note_resolves_to_inter_beside_hand_and_mono():
    from grafli.constants import NOTE_FONT_FAMILY
    view = _view(NOTES)
    faces = {k: QFontInfo(view._note_items[k]._note_font()).family()
             for k in "shm"}
    assert faces == {"s": SANS_FONT_FAMILY, "h": NOTE_FONT_FAMILY,
                     "m": FONT_FAMILY}


def test_a_sans_code_note_stays_mono():
    view = _view('#!grafli v4\n@ note c 0,0 "code: def f(): pass" !sans\n')
    assert view._note_items["c"]._note_font().family() == FONT_FAMILY


def test_set_face_switches_a_note_to_sans_and_back():
    from grafli.constants import NOTE_FONT_FAMILY
    view = _view(NOTES)
    item = view._note_items["h"]
    item.set_face("sans")
    assert item.note.style == "sans"
    assert item._note_font().family() == SANS_FONT_FAMILY
    item.set_face("")
    assert item._note_font().family() == NOTE_FONT_FAMILY


def test_render_draws_a_sans_note_differently_from_handwriting(tmp_path: Path):
    _app()
    one = '#!grafli v4\n@ note n 0,0 "A sans note" !sans\n'
    sans = _render(tmp_path, "sans", one)
    hand = _render(tmp_path, "hand", one.replace(" !sans", "")
                   .replace("v4", "v1"))
    assert sans != hand


# ── !sans in the PDF, PPTX and HTML exports (#206) ──────────────

NOTE_TOUR = ('#!grafli v4\n'
             '@ box a "A" 0,0 160x60\n'
             '@ note n 0,100 "Sans note text" !sans\n'
             '@ bookmark bm "Overlay" @a,n\n'
             '@ bookmark bt "Text" @n\n'
             '@ flow tour "Tour" bm bt\n')


def _pdf_bytes(tmp_path: Path, text: str, name: str) -> bytes:
    from grafli.pdfexport import export_flow_to_pdf
    board = parse(text)
    out = tmp_path / f"{name}.pdf"
    export_flow_to_pdf(_view(text), board.flow_by_id("tour"), out)
    return out.read_bytes()


def test_the_pdf_export_sets_a_sans_note_in_inter(tmp_path: Path):
    assert b"Inter" in _pdf_bytes(tmp_path, NOTE_TOUR, "sans")
    plain = NOTE_TOUR.replace(" !sans", "")
    assert b"Inter" not in _pdf_bytes(tmp_path, plain, "plain")


def test_the_pdf_export_draws_both_note_slides_in_inter(tmp_path: Path,
                                                         monkeypatch):
    from grafli import pdfexport
    families: list[str] = []
    original = pdfexport._draw_markdown

    def spy(painter, rect, body, **kw):
        if "Sans note" in body:
            families.append(kw["family"])
        return original(painter, rect, body, **kw)

    monkeypatch.setattr(pdfexport, "_draw_markdown", spy)
    _pdf_bytes(tmp_path, NOTE_TOUR, "sans")
    assert families == [SANS_FONT_FAMILY, SANS_FONT_FAMILY]
    families.clear()
    _pdf_bytes(tmp_path, NOTE_TOUR.replace(" !sans", ""), "plain")
    assert families == [FONT_FAMILY, FONT_FAMILY]       # as before


def _pptx(tmp_path: Path, text: str, name: str, view=None):
    from pptx import Presentation
    from grafli.pptxexport import export_flow_to_pptx
    board = parse(text)
    out = tmp_path / f"{name}.pptx"
    export_flow_to_pptx(view or _view(text), board.flow_by_id("tour"), out)
    return Presentation(out)


def _note_runs(prs) -> list:
    return [r for slide in list(prs.slides)[1:]
            for sh in slide.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs
            if "Sans note" in r.text]


def test_the_pptx_export_sets_a_sans_note_in_inter(tmp_path: Path):
    runs = _note_runs(_pptx(tmp_path, NOTE_TOUR, "sans"))
    assert len(runs) == 2                               # overlay + text slide
    assert {r.font.name for r in runs} == {SANS_FONT_FAMILY}
    plain = _note_runs(_pptx(tmp_path, NOTE_TOUR.replace(" !sans", ""),
                             "plain"))
    assert {r.font.name for r in plain} == {"+mn-lt"}   # the theme font


def test_the_pptx_export_draws_a_sans_box_in_inter(tmp_path: Path):
    text = ('#!grafli v4\n@ box s "Sans title" 0,0 240x80 %primary !sans\n'
            + TOUR)
    view = _view(text)
    fonts: list[str] = []
    label = view._box_items["s"]._label
    original = label.paint

    def paint(painter, *args):
        fonts.append(QFontInfo(label.font()).family())
        original(painter, *args)

    label.paint = paint
    _pptx(tmp_path, text, "box", view=view)
    assert fonts and set(fonts) == {SANS_FONT_FAMILY}


def _html(tmp_path: Path, text: str, name: str) -> str:
    from grafli.htmlexport import export_html
    _app()
    board = tmp_path / f"{name}.grafli"
    board.write_text(text)
    out = tmp_path / f"{name}.html"
    export_html(board, out)
    return out.read_text()


def _embedded_families(page: str) -> set[str]:
    import json
    import re
    m = re.search(r'id="grafli-fonts">(.*?)</script>', page, re.S)
    return {f["family"] for f in json.loads(m.group(1).replace("<\\/", "</"))}


def test_the_html_export_draws_a_sans_box_and_note_in_inter(tmp_path: Path):
    page = _html(tmp_path, '#!grafli v4\n'
                 '@ box s "Sans title" 0,0 240x80 !sans\n', "box")
    assert 'font-family="Inter"' in page
    assert _embedded_families(page) == {SANS_FONT_FAMILY}
    page = _html(tmp_path, '#!grafli v4\n'
                 '@ note n 0,0 "Sans note text" !sans\n', "note")
    assert 'font-family="Inter"' in page
    assert _embedded_families(page) == {SANS_FONT_FAMILY}


def test_the_html_export_of_a_plain_board_carries_no_inter(tmp_path: Path):
    page = _html(tmp_path, '#!grafli v1\n@ box m "Mono title" 0,0 240x80\n'
                 '@ note n 0,120 "Hand note"\n', "plain")
    assert 'font-family="Inter"' not in page
    assert SANS_FONT_FAMILY not in _embedded_families(page)
