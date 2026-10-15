"""Inline code in a Markdown note renders in the bundled mono face headless,
whatever the note's own face (#219)."""

from __future__ import annotations

import os

import pytest
from PySide6.QtGui import QFontInfo
from PySide6.QtWidgets import QApplication

from grafli.constants import FONT_FAMILY
from grafli.format import parse

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _app() -> QApplication:
    app = QApplication.instance() or QApplication([])
    from grafli.fonts import register_bundled_fonts
    register_bundled_fonts()
    return app


def _code_fragments(doc):
    """The fixed-pitch fragments of a laid-out note document."""
    frags = []
    block = doc.begin()
    while block.isValid():
        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            if frag.isValid() and frag.charFormat().fontFixedPitch():
                frags.append(frag)
            it += 1
        block = block.next()
    return frags


@pytest.mark.parametrize("face", ["", " !sans", " !mono"])
def test_inline_code_in_a_markdown_note_resolves_to_jetbrains_mono(face):
    from grafli.view import GrafliView
    _app()
    view = GrafliView()
    view.load_board(parse(
        '#!grafli v1\n@ note n 0,0 "md: run `grep -n` here"' + face + '\n'))
    doc = view._note_items["n"]._md_document()
    frags = _code_fragments(doc)
    assert [f.text() for f in frags] == ["grep -n"]
    assert QFontInfo(frags[0].charFormat().font()).family() == FONT_FAMILY


CODE_FLOW = '''\
#!grafli v4
@ note n 0,0 """
md:
**Sans face**: `QTextDocument` lays it out
""" !sans
@ bookmark bm "Code" @n
@ flow tour "Tour" bm
'''

_BUNDLED_PDF_FONTS = ("JetBrainsMonoNF-", "Inter-", "PatrickHand-")


def test_inline_code_in_a_pdf_export_embeds_only_the_bundled_faces(tmp_path):
    import re
    from grafli.pdfexport import export_flow_to_pdf
    from grafli.view import GrafliView
    _app()
    board = parse(CODE_FLOW)
    view = GrafliView()
    view.load_board(board)
    out = tmp_path / "tour.pdf"
    export_flow_to_pdf(view, board.flow_by_id("tour"), out)
    fonts = set(re.findall(rb"/BaseFont\s*/([\w+-]+)", out.read_bytes()))
    assert b"JetBrainsMonoNF-Regular" in fonts
    assert all(f.decode().startswith(_BUNDLED_PDF_FONTS) for f in fonts), fonts


def test_inline_code_in_a_pptx_export_names_the_bundled_mono_face(tmp_path):
    from pptx import Presentation
    from grafli.pptxexport import export_flow_to_pptx
    from grafli.view import GrafliView
    _app()
    board = parse(CODE_FLOW)
    view = GrafliView()
    view.load_board(board)
    out = tmp_path / "tour.pptx"
    export_flow_to_pptx(view, board.flow_by_id("tour"), out)
    runs = [r for slide in Presentation(out).slides for sh in slide.shapes
            if sh.has_text_frame for p in sh.text_frame.paragraphs
            for r in p.runs]
    code = [r for r in runs if r.text == "QTextDocument"]
    assert code and all(r.font.name == FONT_FAMILY for r in code)
